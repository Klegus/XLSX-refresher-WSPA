"""
Moodle scanner + heuristic Excel parser for WSPA lesson plans.
Scrapes Moodle "Strefa studenta" for .xlsx files, parses them
without AI, and compares with current plans in MongoDB.
"""
import re
import os
import hashlib
import tempfile
import shutil
from collections import Counter
from datetime import datetime
from urllib.parse import unquote
from difflib import SequenceMatcher

import requests
from bs4 import BeautifulSoup
import openpyxl

from LessonPlanDownloader import LessonPlanDownloader
from shared_utils import get_logger

logger = get_logger('MoodleScanner')

# ─── Moodle scraping ────────────────────────────────────────────────

CATEGORY_URL = "https://puw.wspa.pl/course/index.php?categoryid=1598"

EXCLUDE_PATTERNS = [
    "checklista", "wyrównawcze", "wyrownawcze", "wcag",
    "template", "szablon", "test", "kopia"
]


def get_academic_year_label():
    now = datetime.now()
    # Plany na semestr zimowy pojawiają się na PUW już we wrześniu
    start_year = now.year if now.month >= 9 else now.year - 1
    end_year = start_year + 1
    return f"{start_year % 100}/{end_year % 100}"


def get_current_semester():
    month = datetime.now().month
    return "letni" if 2 <= month <= 8 else "zimowy"


def login_puw(session):
    """Login to PUW using backend credentials (EMAIL/PASSWORD)."""
    url_login = "https://puw.wspa.pl/login/index.php"
    page = session.get(url_login, timeout=30)
    token_match = re.search(r'name="logintoken"\s+value="([^"]+)"', page.text)
    logintoken = token_match.group(1) if token_match else ''

    resp = session.post(url_login, data={
        'anchor': '',
        'logintoken': logintoken,
        'username': os.getenv('EMAIL'),
        'password': os.getenv('PASSWORD')
    }, allow_redirects=True, timeout=30)

    # Strona logowania też zawiera "wyloguj"/"logout", więc sprawdzamy URL:
    # po udanym logowaniu Moodle przekierowuje poza /login/
    return '/login/' not in resp.url


def is_plan_xlsx(filename):
    name_lower = filename.lower()
    if not name_lower.endswith('.xlsx'):
        return False
    for pattern in EXCLUDE_PATTERNS:
        if pattern in name_lower:
            return False
    if 'semestr' not in name_lower:
        return False
    return True


def scrape_strefa_courses(session, url, year_label):
    resp = session.get(url, timeout=30)
    soup = BeautifulSoup(resp.text, 'html.parser')

    courses = []
    subcategory_urls = []

    for link in soup.select('a[href*="course/index.php?categoryid="]'):
        href = link.get('href', '')
        if '&lang=' not in href and href != url:
            subcategory_urls.append(href)

    for link in soup.select('a[href*="course/view.php?id="]'):
        href = link.get('href', '')
        text = link.get_text(strip=True)
        if 'strefa studenta' not in text.lower():
            continue
        # Kursy per kierunek nie mają już roku w nazwie; pomijamy tylko te z innym rokiem
        years_in_name = re.findall(r'\b\d{2}/\d{2}\b', text)
        if not years_in_name or year_label in years_in_name:
            courses.append((text, href))

    return list(set(subcategory_urls)), courses


def _allowed(href):
    """Only links on the university platform become download URLs / get followed."""
    from shared_utils import check_download_url, UnsafeDownload
    try:
        check_download_url(href)
        return True
    except UnsafeDownload:
        logger.warning(f"Skipping link outside the allowed hosts: {href[:120]}")
        return False


def scrape_course_for_xlsx(session, course_url):
    resp = session.get(course_url, timeout=30)
    soup = BeautifulSoup(resp.text, 'html.parser')
    xlsx_links = []

    for link in soup.find_all('a', href=True):
        href = link.get('href', '')
        if '.xlsx' in href.lower() and _allowed(href):
            text = link.get_text(strip=True)
            if is_plan_xlsx(text or unquote(href.split('/')[-1].split('?')[0])):
                xlsx_links.append((text, href))

    for link in soup.find_all('a', href=True):
        href = link.get('href', '')
        if 'mod/folder/view.php' in href and _allowed(href):
            folder_resp = session.get(href, timeout=30)
            folder_soup = BeautifulSoup(folder_resp.text, 'html.parser')
            for flink in folder_soup.find_all('a', href=True):
                fhref = flink.get('href', '')
                fname = flink.get_text(strip=True)
                if '.xlsx' in fhref.lower() and _allowed(fhref):
                    display = fname if fname else unquote(fhref.split('/')[-1].split('?')[0])
                    if is_plan_xlsx(display):
                        xlsx_links.append((display, fhref))
    return xlsx_links


def scrape_all_xlsx_links(session, progress_cb=None):
    year_label = get_academic_year_label()
    semester = get_current_semester()

    if progress_cb:
        progress_cb(f"Szukam kursów Strefa studenta ({year_label})...")

    all_courses = []
    subcats, courses = scrape_strefa_courses(session, CATEGORY_URL, year_label)
    all_courses.extend(courses)

    for subcat_url in set(subcats):
        sub_subcats, sub_courses = scrape_strefa_courses(session, subcat_url, year_label)
        all_courses.extend(sub_courses)
        for sub2_url in set(sub_subcats):
            _, sub2_courses = scrape_strefa_courses(session, sub2_url, year_label)
            all_courses.extend(sub2_courses)

    # Deduplicate courses
    seen = set()
    unique_courses = []
    for name, url in all_courses:
        if url not in seen:
            seen.add(url)
            unique_courses.append((name, url))

    if progress_cb:
        progress_cb(f"Znaleziono {len(unique_courses)} kursów, szukam .xlsx...")

    all_xlsx = []
    for i, (course_name, course_url) in enumerate(unique_courses):
        xlsx = scrape_course_for_xlsx(session, course_url)
        for name, link in xlsx:
            if semester in name.lower():
                all_xlsx.append((name, link))

    # Deduplicate by URL
    seen_links = set()
    unique_xlsx = []
    for name, link in all_xlsx:
        if link not in seen_links:
            seen_links.add(link)
            unique_xlsx.append((name, link))

    if progress_cb:
        progress_cb(f"Znaleziono {len(unique_xlsx)} plików .xlsx ({year_label}, sem. {semester})")

    return unique_xlsx, year_label, semester


# ─── Heuristic Excel parser ────────────────────────────────────────

DAYS_CANONICAL = {
    'PONIEDZIAŁEK': 'PONIEDZIAŁEK', 'PONIEDZIALEK': 'PONIEDZIAŁEK',
    'WTOREK': 'WTOREK',
    'ŚRODA': 'ŚRODA', 'SRODA': 'ŚRODA',
    'CZWARTEK': 'CZWARTEK',
    'PIĄTEK': 'PIĄTEK', 'PIATEK': 'PIĄTEK',
    'SOBOTA': 'SOBOTA',
    'NIEDZIELA': 'NIEDZIELA',
}

ROMAN_NUMERALS = {'I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X'}


def fuzzy_match_day(text):
    text = text.upper().strip()
    if text in DAYS_CANONICAL:
        return DAYS_CANONICAL[text]
    best_match, best_score = None, 0
    for variant, canonical in DAYS_CANONICAL.items():
        score = SequenceMatcher(None, text, variant).ratio()
        if score > best_score:
            best_score = score
            best_match = canonical
    return best_match if best_score >= 0.8 else None


def find_days_row(ws, max_search=6):
    for row_idx in range(1, max_search + 1):
        days_found = {}
        for col in range(1, min(ws.max_column + 1, 40)):
            v = ws.cell(row_idx, col).value
            if v and isinstance(v, str):
                day = fuzzy_match_day(v.strip())
                if day:
                    days_found[col] = day
        if len(days_found) >= 2:
            return row_idx, days_found
    return None, {}


def find_groups_row(ws, days_row_idx):
    if days_row_idx is None:
        return None
    candidate = days_row_idx + 1
    if candidate > ws.max_row:
        return None
    col1_val = ws.cell(candidate, 1).value
    if col1_val and isinstance(col1_val, str):
        if 'GRUPA' in col1_val.upper().strip() or 'GRUPY' in col1_val.upper().strip():
            return candidate
    return None


def extract_groups(ws, groups_row_idx):
    if groups_row_idx is None:
        return {}
    raw_groups = {}
    for col in range(2, ws.max_column + 1):
        v = ws.cell(groups_row_idx, col).value
        if v and isinstance(v, str) and v.strip():
            raw_groups[col] = v
    seen_values = {}
    for col, original in raw_groups.items():
        clean = re.sub(r'\s+', ' ', original.replace('\n', ' ')).strip()
        clean = clean.rstrip(':').rstrip('.').strip()
        normalized = clean.lower()
        if normalized not in seen_values:
            seen_values[normalized] = (clean, original)
    return {clean: original for clean, original in seen_values.values()}


def _squash(text):
    import unicodedata
    t = unicodedata.normalize('NFKD', str(text or '').lower().replace('ł', 'l'))
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in t if not unicodedata.combining(c)))


def _is_variant(short, long):
    """'Grupa 1' -> 'Grupa 1 podział wg nazwisk' yes; 'Grupa 1' -> 'Grupa 10' no."""
    a, b = _squash(short), _squash(long)
    if a == b:
        return True
    return b.startswith(a) and not b[len(a)].isdigit()


def merge_group_variants(ws, groups_row_idx, groups, days_found):
    """Merge headers that name the same group differently on different days.

    Excel files are edited by hand, so one group can be "Grupy 1-18" on Thursday
    and "Grupy 1 - 18" on Sunday, or "Grupa 1" on Friday and "Grupa 1 podział wg
    nazwisk: A-D" on Saturday. Two headers are treated as one group when they are
    variants of each other and never appear on the same day (groups like
    "Sp.: X" and "Sp.: X gr.1" that sit side by side stay separate).
    Returns (groups, aliases) where aliases maps group name -> other header texts.
    """
    if groups_row_idx is None or len(groups) < 2:
        return groups, {}
    day_starts = sorted(days_found)

    def day_of(col):
        starts = [c for c in day_starts if c <= col]
        return days_found[starts[-1]] if starts else None

    days_of = {clean: set() for clean in groups}
    original_to_clean = {orig: clean for clean, orig in groups.items()}
    for col in range(2, ws.max_column + 1):
        clean = original_to_clean.get(ws.cell(groups_row_idx, col).value)
        if clean:
            days_of[clean].add(day_of(col))

    # A shorter header that is a variant of longer ones on other days is either
    # the same group written differently (one match -> merge) or a shared header
    # for several sub-groups, e.g. "Sp.: X" on days with joint classes and
    # "Sp.: X gr.1" / "Sp.: X gr.2" on the others (many matches -> every
    # sub-group also gets the shared columns)
    names = sorted(groups, key=len, reverse=True)
    merged_into, aliases = {}, {}
    for short_name in sorted(groups, key=len):
        targets = [long_name for long_name in names
                   if long_name != short_name and long_name not in merged_into
                   and len(long_name) >= len(short_name)
                   and _is_variant(short_name, long_name)
                   and not (days_of[short_name] & days_of[long_name])]
        if not targets:
            continue
        merged_into[short_name] = targets
        for long_name in targets:
            aliases.setdefault(long_name, []).append(groups[short_name])

    kept = {name: orig for name, orig in groups.items() if name not in merged_into}
    if merged_into:
        logger.info(f"Merged group header variants: {merged_into}")
    return kept, aliases


def count_group_columns(ws, groups_row_idx, groups, aliases=None):
    if groups_row_idx is None:
        return {}
    original_to_clean = {orig: clean for clean, orig in groups.items()}
    for clean, variants in (aliases or {}).items():
        for variant in variants:
            original_to_clean[variant] = clean
    counts = Counter()
    for col in range(2, ws.max_column + 1):
        v = ws.cell(groups_row_idx, col).value
        if v and v in original_to_clean:
            counts[original_to_clean[v]] += 1
    return dict(counts)


def detect_category(days_found):
    day_set = set(days_found.values())
    if day_set == {'SOBOTA', 'NIEDZIELA'}:
        return 'nst-online'
    if day_set & {'PONIEDZIAŁEK', 'WTOREK', 'ŚRODA', 'CZWARTEK'}:
        return 'st'
    if 'PIĄTEK' in day_set and ('SOBOTA' in day_set or 'NIEDZIELA' in day_set):
        return 'nst'
    if day_set <= {'SOBOTA', 'NIEDZIELA'}:
        return 'nst-online'
    return 'st'


def extract_faculty_from_header(ws):
    val = ws.cell(1, 1).value
    if not val or not isinstance(val, str):
        return None
    header = val.split('\n')[0].strip()
    words = header.split()
    semestr_idx = None
    for i, w in enumerate(words):
        if w.upper() in ('SEMESTR', 'SEM.', 'SEM'):
            semestr_idx = i
            break
    if semestr_idx is None:
        faculty_words = []
        for word in words:
            if word.isdigit():
                break
            faculty_words.append(word)
        return ' '.join(faculty_words).strip(' -') if faculty_words else None
    cut_idx = semestr_idx
    for i in range(max(0, semestr_idx - 2), semestr_idx):
        if words[i].upper().rstrip(',').rstrip('.') in ROMAN_NUMERALS:
            cut_idx = i
            break
    faculty_words = words[:cut_idx]
    return ' '.join(faculty_words).strip(' -') if faculty_words else None


def extract_faculty_from_filename(filename):
    name = unquote(filename).rsplit('.', 1)[0]
    parts = name.split(' - ')
    return parts[0].strip() if parts else name


def is_schedule_sheet(ws):
    row, _ = find_days_row(ws)
    return row is not None


def detect_mixed_plan(col_counts, category):
    if not col_counts or len(col_counts) < 2:
        return False
    expected = {'st': 5, 'nst': 3, 'nst-online': 2}.get(category, 5)
    has_partial = any(0 < c < expected for c in col_counts.values())
    has_full = any(c >= expected for c in col_counts.values())
    if has_partial and has_full:
        return True
    unique_counts = set(col_counts.values())
    return len(unique_counts) > 1 and 0 not in unique_counts


def sanitize_string(text):
    if not text:
        return "unknown"
    text = text.lower().strip().replace('.', '_')
    text = "".join(c if c.isalnum() or c == '_' else '_' for c in text)
    while '__' in text:
        text = text.replace('__', '_')
    return text.strip('_') or "unknown"


def generate_plan_key(faculty, sheet_name, existing_keys):
    base = f"{sanitize_string(faculty)}_{sanitize_string(sheet_name)}"
    key = base
    counter = 1
    while key in existing_keys:
        key = f"{base}_{counter}"
        counter += 1
    return key


def process_excel_file(filepath, download_url, filename):
    results = []
    try:
        wb = openpyxl.load_workbook(filepath)
    except Exception as e:
        logger.error(f"Error loading {filename}: {e}")
        return results

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if not is_schedule_sheet(ws):
            continue

        days_row_idx, days_found = find_days_row(ws)
        groups_row_idx = find_groups_row(ws, days_row_idx)
        groups = extract_groups(ws, groups_row_idx)
        groups, aliases = merge_group_variants(ws, groups_row_idx, groups, days_found)
        col_counts = count_group_columns(ws, groups_row_idx, groups, aliases)
        category = detect_category(days_found)

        faculty = extract_faculty_from_header(ws)
        if not faculty:
            faculty = extract_faculty_from_filename(filename)

        file_base = filename.rsplit('.', 1)[0]
        nice_name = f"{file_base} - {sheet_name.strip()}".replace('.', '').strip()

        is_mixed = detect_mixed_plan(col_counts, category)

        plan = {
            'name': nice_name,
            'faculty': faculty.lower(),
            'groups': groups if groups else {},
            'category': category,
            'download_url': download_url,
            'sheet_name': sheet_name,
        }
        if aliases:
            plan['group_aliases'] = aliases
        if is_mixed:
            plan['mixed'] = True
            plan['groups_column_info'] = col_counts

        results.append(plan)

    wb.close()
    return results


# ─── Full scan orchestrator ─────────────────────────────────────────

def run_full_scan(progress_cb=None):
    """
    Run the complete scan: scrape Moodle -> download Excels -> parse -> return plans dict.
    progress_cb(message: str) is called with status updates.
    """
    session = requests.Session()
    session.headers.update({
        'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36'
    })

    if progress_cb:
        progress_cb("Logowanie do PUW...")

    if not login_puw(session):
        raise RuntimeError("Login do PUW nieudany. Sprawdź EMAIL/PASSWORD.")

    xlsx_links, year_label, semester = scrape_all_xlsx_links(session, progress_cb)

    if not xlsx_links:
        return {}, year_label, semester

    username = os.getenv('EMAIL')
    password = os.getenv('PASSWORD')
    tmp_dir = tempfile.mkdtemp(prefix='moodle_scan_')
    all_plans = {}

    try:
        for i, (name, url) in enumerate(xlsx_links):
            if progress_cb:
                progress_cb(f"Pobieranie i parsowanie {i+1}/{len(xlsx_links)}: {name}")

            url_hash = hashlib.md5(url.encode(), usedforsecurity=False).hexdigest()[:12]
            local_path = os.path.join(tmp_dir, f'{url_hash}.xlsx')

            try:
                downloader = LessonPlanDownloader(username, password, tmp_dir, url)
                downloader.download_file()
                src = downloader.get_file_save_path()
                if not src:
                    logger.warning(f"Download failed for {name}")
                    continue
                shutil.copy(src, local_path)

                filename = unquote(url.split('/')[-1].split('?')[0])
                results = process_excel_file(local_path, url, filename)

                for plan in results:
                    key = generate_plan_key(plan['faculty'], plan['sheet_name'], all_plans)
                    all_plans[key] = plan

            except Exception as e:
                logger.error(f"Error processing {name}: {e}")
                continue
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    if progress_cb:
        progress_cb(f"Parsowanie zakończone: {len(all_plans)} planów")

    return all_plans, year_label, semester


def compare_plans(scraped, current):
    """
    Compare scraped plans with current MongoDB plans.
    Returns dict with new, changed, removed.
    """
    scraped_keys = set(scraped.keys())
    current_keys = set(current.keys())

    new_plans = {k: scraped[k] for k in scraped_keys - current_keys}
    removed_plans = {k: current[k] for k in current_keys - scraped_keys}

    changed_plans = {}
    compare_fields = ['download_url', 'sheet_name', 'category', 'faculty', 'group_aliases', 'mixed']

    for key in scraped_keys & current_keys:
        diff_fields = []
        for field in compare_fields:
            if scraped[key].get(field) != current[key].get(field):
                diff_fields.append(field)

        # Compare groups - normalize empty variants (null, {}, {"cały kierunek": "all"})
        s_groups = scraped[key].get('groups') or {}
        c_groups = current[key].get('groups') or {}
        if c_groups == {"cały kierunek": "all"}:
            c_groups = {}
        if s_groups != c_groups:
            diff_fields.append('groups')

        if diff_fields:
            changed_plans[key] = {
                'old': current[key],
                'new': scraped[key],
                'diff_fields': diff_fields
            }

    return {
        'new': new_plans,
        'changed': changed_plans,
        'removed': removed_plans,
    }


# ─── Automatic scan (called from the backend loop) ─────────────────

def auto_scan(db):
    """Scan PUW and add plans that appeared since the last scan.

    New plans are added right away (source validation still quarantines any
    that do not parse cleanly). Changed and removed plans are only reported -
    they can drop or rename groups students already use, so an admin approves
    them in the panel. Returns the summary stored in system_config.auto_scan.
    """
    started = datetime.now()
    summary = {"at": started, "added": [], "pending_changed": [], "pending_removed": [], "error": None}
    try:
        scraped, year_label, semester = run_full_scan()
        config = db.plans_config.find_one({"_id": "plans_json"}) or {}
        current = config.get("plans") or {}
        diff = compare_plans(scraped, current)

        if diff["new"]:
            current.update(diff["new"])
            db.plans_config.update_one(
                {"_id": "plans_json"},
                {"$set": {"plans": current, "last_updated": started.isoformat(), "source": "auto_scan"}},
                upsert=True)
            logger.info(f"Auto-scan added {len(diff['new'])} new plans")

        summary.update(
            year=year_label, semester=semester, scraped=len(scraped),
            added=[p.get("name", k) for k, p in diff["new"].items()],
            pending_changed=sorted(diff["changed"]),
            pending_removed=sorted(diff["removed"]),
        )
    except Exception as e:
        logger.error(f"Auto-scan failed: {e}")
        summary["error"] = str(e)
    summary["duration"] = round((datetime.now() - started).total_seconds(), 1)
    db.system_config.update_one({"_id": "config"}, {"$set": {"auto_scan": summary}}, upsert=True)
    return summary
